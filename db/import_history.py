"""db/import_history.py — Parse and import historic XC results from xlsx spreadsheets."""

import re
from datetime import datetime, timedelta, time as dt_time

import openpyxl

from db.connection import get_connection, release_connection, fetchone
from db.season import get_or_create_season
from db.meet import add_meet
from db.xc import save_xc_result, recalculate_xc_pr_flags
from db.athlete import add_athlete, add_to_roster


def _cell_to_seconds(value) -> float | None:
    """Convert an xlsx cell value to XC finish time in seconds.

    Handles timedelta and datetime.time objects from openpyxl.
    Reinterprets times >1h as Excel HH:MM:SS misreadings of MM:SS.
    Skips negative timedeltas (improvement calculations) and strings.
    """
    if value is None:
        return None

    if isinstance(value, str):
        return None

    if isinstance(value, timedelta):
        total = value.total_seconds()
        if total <= 0:
            return None
        if total > 3600:
            h = int(total // 3600)
            m = int((total % 3600) // 60)
            s = total % 60
            return h * 60 + m + s / 100
        return total

    if isinstance(value, dt_time):
        total = value.hour * 3600 + value.minute * 60 + value.second + value.microsecond / 1_000_000
        if total > 3600:
            h = int(total // 3600)
            m = int((total % 3600) // 60)
            s = total % 60
            return h * 60 + m + s / 100
        return total

    return None


def _seconds_to_time_str(seconds: float) -> str:
    """Format seconds as M:SS.ss for storage."""
    m = int(seconds // 60)
    s = seconds % 60
    return f"{m}:{s:05.2f}"


def _parse_date(value, year: int) -> str | None:
    """Parse a date from row 1 — datetime object or M/D string."""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, str):
        match = re.match(r"^(\d{1,2})/(\d{1,2})$", value.strip())
        if match:
            month, day = int(match.group(1)), int(match.group(2))
            return f"{year}-{month:02d}-{day:02d}"
    return None


def _clean_venue(venue: str) -> tuple[str, str, bool]:
    """Clean venue name. Returns (name, distance, is_cancelled)."""
    venue = venue.strip()
    is_cancelled = bool(re.search(r"cancel+ed", venue, re.IGNORECASE))

    distance = "2mi"
    if re.search(r"\(3k\)", venue, re.IGNORECASE):
        distance = "3K"

    clean = re.sub(r"\s*\([^)]*\)\s*", " ", venue).strip()
    clean = re.sub(r"\s*Cancel+ed\s*$", "", clean, flags=re.IGNORECASE).strip()
    return clean, distance, is_cancelled


_STATS_LABELS = {"AVERAGE", "MEDIAN", "1-5 SPLIT", "TOP5 AVG"}


def parse_xc_xlsx(filepath: str, year: int) -> dict:
    """Parse an XC spreadsheet into structured meet and athlete data.

    Returns {"meets": [...], "athletes": [...]}.
    """
    wb = openpyxl.load_workbook(filepath, data_only=True)

    all_meets: list[dict] = []
    all_athletes: list[dict] = []

    for sheet_name in ("BOYS", "GIRLS"):
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        gender = "M" if sheet_name == "BOYS" else "F"

        meet_cols: list[dict] = []
        for col in range(1, ws.max_column + 1):
            r2 = ws.cell(2, col).value
            if not r2 or not str(r2).strip():
                continue

            date_val = ws.cell(1, col).value
            meet_date = _parse_date(date_val, year)
            if not meet_date:
                continue

            venue, distance, is_cancelled = _clean_venue(str(r2))
            if is_cancelled:
                continue

            meet_cols.append({
                "col": col,
                "date": meet_date,
                "venue": venue,
                "distance": distance,
            })

        col_to_meet_idx: dict[int, int] = {}
        for mc in meet_cols:
            found_idx = None
            for i, existing in enumerate(all_meets):
                if existing["date"] == mc["date"]:
                    found_idx = i
                    break
            if found_idx is not None:
                col_to_meet_idx[mc["col"]] = found_idx
            else:
                col_to_meet_idx[mc["col"]] = len(all_meets)
                all_meets.append({
                    "date": mc["date"],
                    "venue": mc["venue"],
                    "distance": mc["distance"],
                })

        for row in range(3, ws.max_row + 1):
            last_name = ws.cell(row, 1).value
            first_name = ws.cell(row, 2).value

            if not last_name or not first_name:
                continue
            if str(last_name).strip().upper() in _STATS_LABELS:
                continue
            if str(first_name).strip().upper() in _STATS_LABELS:
                continue

            last_name = str(last_name).strip()
            first_name = str(first_name).strip()

            results = []
            for mc in meet_cols:
                cell_val = ws.cell(row, mc["col"]).value
                seconds = _cell_to_seconds(cell_val)
                if seconds and 300 < seconds < 2400:
                    results.append({
                        "meet_idx": col_to_meet_idx[mc["col"]],
                        "time_seconds": seconds,
                        "time_formatted": _seconds_to_time_str(seconds),
                    })

            if results:
                all_athletes.append({
                    "first": first_name,
                    "last": last_name,
                    "gender": gender,
                    "results": results,
                })

    return {"meets": all_meets, "athletes": all_athletes}


def import_xc_season(parsed: dict, school_id: int, year: int) -> dict:
    """Import parsed XC data into the database. Returns summary counts."""
    season_id = get_or_create_season(year, "XC", school_id)

    used_meet_idxs = {
        r["meet_idx"]
        for a in parsed["athletes"]
        for r in a["results"]
    }

    meet_ids: dict[int, int] = {}
    meets_created = 0
    for idx, meet_info in enumerate(parsed["meets"]):
        if idx not in used_meet_idxs:
            continue

        conn = get_connection()
        try:
            existing = fetchone(
                conn,
                "SELECT id FROM meet WHERE season_id=? AND meet_date=?",
                (season_id, meet_info["date"]),
            )
        finally:
            release_connection(conn)

        if existing:
            meet_ids[idx] = existing["id"]
        else:
            mid = add_meet(
                season_id,
                meet_info["venue"],
                meet_info["date"],
                meet_info["venue"],
                school_id,
            )
            meet_ids[idx] = mid
            meets_created += 1

    athletes_created = 0
    athletes_matched = 0
    results_imported = 0

    for athlete in parsed["athletes"]:
        conn = get_connection()
        try:
            existing = fetchone(
                conn,
                """SELECT id FROM athlete
                   WHERE LOWER(first_name)=LOWER(?)
                     AND LOWER(last_name)=LOWER(?)
                     AND school_id=?""",
                (athlete["first"], athlete["last"], school_id),
            )
        finally:
            release_connection(conn)

        if existing:
            athlete_id = existing["id"]
            athletes_matched += 1
        else:
            athlete_id = add_athlete(
                athlete["first"],
                athlete["last"],
                8,
                athlete["gender"],
                school_id,
            )
            athletes_created += 1

        add_to_roster(season_id, athlete_id)

        for result in athlete["results"]:
            meet_idx = result["meet_idx"]
            if meet_idx not in meet_ids:
                continue
            meet_id = meet_ids[meet_idx]
            meet_info = parsed["meets"][meet_idx]
            save_xc_result(
                meet_id=meet_id,
                athlete_id=athlete_id,
                finish_time=result["time_formatted"],
                distance=meet_info["distance"],
            )
            results_imported += 1

    recalculate_xc_pr_flags()

    return {
        "season_id": season_id,
        "meets_created": meets_created,
        "athletes_created": athletes_created,
        "athletes_matched": athletes_matched,
        "results_imported": results_imported,
    }
